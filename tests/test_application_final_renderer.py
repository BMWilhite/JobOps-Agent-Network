"""Entirely fictional offline regressions for heading and pipe-title formats."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from application_documents import DocumentBlock, ResumeSection, SourceApplicationDocuments
from application_rendering import catalog_for_documents
from generate_application import validate_documents, write_package

PROFILE = """# Fictional profile: invented parcel-workshop scenario
## Experience
### Fable Parcel Workshop (Fictional)
Dispatch Supervisor | 2016-2019

- ISO 9001 document filing.
- Generic parcel service responsibilities.
- Redesigned warehouse bin labels to reduce routing errors.
- Identified bottlenecks in parcel dispatch reviews.
### Imaginary Training Cooperative (Fictional)
Dispatch Coordinator / Training Coordinator
Regional Service Enablement | 2010-2014

- Approximately 40% onboarding-time improvement.
- Reduced onboarding time from approximately 30 days to 18 days.
- Identified packing improvements worth approximately 8% additional monthly capacity.
- Do not invent revenue or growth accomplishments.
## Recent Technical Project: Museum Inventory Demo
Built a local museum inventory prototype with SQLite and substantial AI-assisted coding.
"""
POSTING = "Company: Imaginary Parcel Labs (Fictional)\nJob title: Warehouse Process Analyst\nFull job description: Improve onboarding systems and parcel workflows."


def selection(profile=PROFILE):
    entries = catalog_for_documents(profile)
    ids = [e['evidence_id'] for e in entries if 'Project' not in e['section']]
    letter = next(e['evidence_id'] for e in entries if 'Reduced onboarding' in e['text'])
    return SourceApplicationDocuments(
        resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=ids, phrasing='source')])],
        cover_letter_paragraphs=[DocumentBlock(evidence_ids=[letter], phrasing='first_person')],
        verification_items=[])


class FinalRendererTests(unittest.TestCase):
    def setUp(self):
        self.block = patch('agents.Runner.run_sync', side_effect=AssertionError('No paid calls')).start()
        self.addCleanup(patch.stopall)

    def test_heading_pipe_title_dates_and_real_docx(self):
        docs = validate_documents(selection(), PROFILE, POSTING)
        with tempfile.TemporaryDirectory() as root:
            folder = write_package(docs, '', output_root=Path(root))
            word = Document(folder / 'Resume.docx')
            headers = [p.text for p in word.paragraphs if p.style.name == 'Heading 3']
            self.assertEqual(headers, ['Fable Parcel Workshop (Fictional)', 'Imaginary Training Cooperative (Fictional)'])
            text = [p.text for p in word.paragraphs]
            self.assertIn('Dispatch Supervisor | 2016-2019', text)
            self.assertIn('Dispatch Coordinator / Training Coordinator Regional Service Enablement | 2010-2014', text)
            bullets = [p.text for p in word.paragraphs if p.style.name == 'List Bullet']
            self.assertFalse(any('2016-2019' in b or '2010-2014' in b or 'Coordinator /' in b for b in bullets))
            self.assertGreaterEqual(len(bullets), 3)
            self.assertIn('[Applicant name', docs.resume_markdown)
            self.assertIn('[Email / phone', docs.resume_markdown)
            self.assertNotIn('Do not invent', docs.resume_markdown)
            letter = Document(folder / 'Cover_Letter.docx')
            self.assertTrue(any('Imaginary Parcel Labs (Fictional)' in p.text for p in letter.paragraphs))

    def test_header_provenance_resolves_multiline_title(self):
        docs = validate_documents(selection(), PROFILE, POSTING)
        entry = next(e for e in catalog_for_documents(PROFILE) if 'Dispatch Coordinator /' in e['text'])
        ref = next(e for e in docs.evidence if e.claim.startswith('Dispatch Coordinator /'))
        self.assertEqual(ref.evidence_id, entry['evidence_id'])
        self.assertEqual(ref.source_quote, entry['text'])
        self.assertEqual(PROFILE[ref.source_start:ref.source_end], ref.source_quote)

    def test_duplicate_percentage_removed_in_either_selection_order(self):
        for reverse in (False, True):
            draft = selection()
            if reverse:
                draft.resume_sections[0].blocks[0].evidence_ids.reverse()
            docs = validate_documents(draft, PROFILE, POSTING)
            self.assertIn('approximately 30 days to 18 days', docs.resume_markdown)
            self.assertNotIn('40%', docs.resume_markdown)
            self.assertIn('approximately 8% additional monthly capacity', docs.resume_markdown)

    def test_percentage_alone_is_retained(self):
        draft = selection()
        duration = next(e['evidence_id'] for e in catalog_for_documents(PROFILE) if 'Reduced onboarding' in e['text'])
        draft.resume_sections[0].blocks[0].evidence_ids.remove(duration)
        self.assertIn('40%', validate_documents(draft, PROFILE, POSTING).resume_markdown)

    def test_concrete_actions_rank_before_generic_numeric_responsibilities(self):
        docs = validate_documents(selection(), PROFILE, POSTING)
        self.assertLess(docs.resume_markdown.index('Redesigned'), docs.resume_markdown.index('ISO 9001'))
        self.assertLess(docs.resume_markdown.index('Identified bottlenecks'), docs.resume_markdown.index('Generic parcel'))

    def test_project_is_labeled_and_cannot_be_employment(self):
        eid = next(e['evidence_id'] for e in catalog_for_documents(PROFILE) if 'Project' in e['section'])
        draft = selection()
        draft.resume_sections.append(ResumeSection(heading='Projects', blocks=[DocumentBlock(evidence_ids=[eid], phrasing='source')]))
        docs = validate_documents(draft, PROFILE, POSTING)
        self.assertIn('## Projects', docs.resume_markdown)
        self.assertIn('### Recent Technical Project: Museum Inventory Demo', docs.resume_markdown)
        self.assertIn('substantial AI-assisted coding', docs.resume_markdown)
        draft.resume_sections[-1].heading = 'Experience'
        with self.assertRaisesRegex(ValueError, 'project evidence'):
            validate_documents(draft, PROFILE, POSTING)

    def test_unknown_employer_is_not_inferred_from_generic_heading(self):
        profile = PROFILE.replace('### Fable Parcel Workshop (Fictional)', '### Experience')
        docs = validate_documents(selection(profile), profile, POSTING)
        self.assertIn('[Employer', docs.resume_markdown)
        self.assertNotIn('### Experience', docs.resume_markdown)

    def test_duplicate_metric_in_distinct_jobs_is_retained(self):
        profile = PROFILE + '\n## Experience\n### Imaginary Archive (Fictional)\nRecords Assistant | 2005-2008\n\n- Approximately 40% onboarding-time improvement.\n'
        docs = validate_documents(selection(profile), profile, POSTING)
        self.assertEqual(docs.resume_markdown.count('40%'), 1)

    def test_plain_limitations_label_excludes_following_metadata(self):
        profile = PROFILE.replace('### Imaginary Training Cooperative (Fictional)', '\nImportant limitations:\n\n- Internal note: pilot workshop only.\n- Do not invent certifications.\n### Imaginary Training Cooperative (Fictional)')
        docs = validate_documents(selection(profile), profile, POSTING)
        self.assertNotIn('pilot workshop only', docs.resume_markdown)
        self.assertNotIn('Important limitations', docs.resume_markdown)
        self.assertIn('approximately 30 days to 18 days', docs.resume_markdown)
