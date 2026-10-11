"""Fictional quality fixtures exercise source resolution and real Word structure."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

from application_documents import DocumentBlock, ResumeSection, SourceApplicationDocuments
from application_evidence import document_catalog
from application_rendering import validate_quality
from generate_application import validate_documents, write_package


PROFILE = """# Fictional candidate profile
## Contact
- **Name:** Jordan Example
- Email: jordan@example.test
- Phone: +1 555 010 1234
## Experience
### Example Services
- Company: Example Services
- Actual job title: Implementation Coordinator
- Dates: 2021-2023
- Coordinated onboarding for 12 employees. Built a workflow tracker. Trained colleagues on documented procedures. Maintained operational reports. Supported system rollouts. Documented insurance verification workflows. Improved handoff checklists. Organized weekly reviews.
- Identified $10,000 in annual savings potential (estimated). Do not present estimated improvements as audited realized savings.
- Do not invent revenue, headcount, or growth accomplishments.
## Skills
- Moderate Excel; beginner SQL; AI-assisted drafting.
## Internal instructions
- Never invent fintech experience. Follow the prompt guidance.
"""
POSTING = """Company: Stedi
Job title: Business Operations & Systems
Full job description:
Support operational workflows, implementation, systems, onboarding and process improvements.
"""


def fixture():
    catalog = document_catalog(PROFILE)
    merged = next(e['evidence_id'] for e in catalog if 'Coordinated onboarding' in e['text'])
    savings = next(e['evidence_id'] for e in catalog if 'Identified $10,000' in e['text'])
    instruction = next(e['evidence_id'] for e in catalog if 'Do not invent revenue' in e['text'])
    return SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[
        DocumentBlock(evidence_ids=[merged, savings, instruction], phrasing='source')])],
        cover_letter_paragraphs=[DocumentBlock(evidence_ids=[merged], unit_indices=[1, 2], phrasing='first_person'),
            DocumentBlock(evidence_ids=[savings], unit_indices=[1], phrasing='first_person')], verification_items=[])


class ApplicationQualityTests(unittest.TestCase):
    def setUp(self):
        self.runner = patch('agents.Runner.run_sync', side_effect=AssertionError('Paid calls forbidden')).start()
        self.addCleanup(patch.stopall)

    def resolve(self, draft=None, profile=PROFILE, posting=POSTING):
        return validate_documents(draft or fixture(), profile, posting)

    def test_docx_discrete_bullets_job_fields_and_ats_structure(self):
        docs = self.resolve()
        with tempfile.TemporaryDirectory() as temporary:
            folder = write_package(docs, 'https://example.test/apply', output_root=Path(temporary))
            word = Document(folder / 'Resume.docx')
            bullets = [p.text for p in word.paragraphs if p.style.name == 'List Bullet']
            self.assertGreaterEqual(len(bullets), 3)
            self.assertLessEqual(len(bullets), 5)
            self.assertTrue(all(len(b.split()) <= 45 for b in bullets))
            self.assertTrue(any(p.text == 'Experience' and p.style.name == 'Heading 2' for p in word.paragraphs))
            self.assertTrue(any(p.text == 'Example Services' and p.style.name == 'Heading 3' for p in word.paragraphs))
            self.assertIn('Implementation Coordinator | 2021-2023', [p.text for p in word.paragraphs])
            self.assertEqual(word.paragraphs[0].text, 'Jordan Example')
            self.assertFalse(word.tables)
            provenance = json.loads((folder / 'source_evidence.json').read_text())
            catalog = {e['evidence_id']: e for e in document_catalog(PROFILE)}
            for item in provenance['evidence']:
                if item['source'] == 'job_posting':
                    self.assertEqual(item['source_quote'], POSTING)
                    continue
                self.assertEqual(item['source_quote'], catalog[item['evidence_id']]['text'])
                self.assertIn(item['claim'], getattr(docs, item['document'] + '_markdown'))
            self.assertTrue(any(item['source_unit_index'] for item in provenance['evidence']))
        self.runner.assert_not_called()

    def test_instructions_and_limitations_excluded_even_when_selected(self):
        docs = self.resolve()
        for text in [docs.resume_markdown, docs.cover_letter_markdown]:
            self.assertNotIn('Do not', text)
            self.assertNotIn('prompt guidance', text)
            self.assertNotIn('audited realized', text)
            self.assertNotRegex(text, r'\bP\d{3}\b')
        self.assertTrue(any('excluded' in item for item in docs.verification_items))

    def test_company_specific_letter_is_prose_with_supported_examples(self):
        docs = self.resolve()
        letter = docs.cover_letter_markdown
        self.assertIn('Business Operations & Systems position at Stedi', letter)
        self.assertIn('I coordinated onboarding for 12 employees.', letter)
        self.assertIn('I built a workflow tracker.', letter)
        self.assertIn('I identified $10,000 in annual savings potential (estimated).', letter)
        self.assertIn('support Stedi in this role', letter)
        self.assertIn('Your posting emphasizes operational workflows and process improvements.', letter)
        job_evidence = next(e for e in docs.evidence if e.source == 'job_posting')
        self.assertEqual(job_evidence.source_quote, POSTING)
        self.assertNotIn('- Coordinated', letter)
        self.assertNotIn('Implemented fintech', letter)
        other = self.resolve(posting=POSTING.replace('Stedi', 'Example Health').replace('Business Operations & Systems', 'Implementation Lead'))
        self.assertIn('Implementation Lead position at Example Health', other.cover_letter_markdown)
        self.assertNotIn('Stedi', other.cover_letter_markdown)

    def test_actual_contact_is_sourced_and_missing_contact_has_placeholders(self):
        docs = self.resolve()
        self.assertIn('# Jordan Example', docs.resume_markdown)
        self.assertIn('jordan@example.test', docs.resume_markdown)
        self.assertIn('+1 555 010 1234', docs.cover_letter_markdown)
        profile = '## Experience\n- Coordinated onboarding and training.\n'
        value = SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=['P001'], phrasing='source')])],
            cover_letter_paragraphs=[DocumentBlock(evidence_ids=['P001'], phrasing='first_person')], verification_items=[])
        missing = self.resolve(value, profile)
        self.assertIn('[Applicant name — verify]', missing.resume_markdown)
        self.assertIn('[Email / phone — verify]', missing.cover_letter_markdown)
        self.assertTrue(any('applicant-name' in item for item in missing.verification_items))
        self.assertNotIn('jordan@example.test', missing.resume_markdown)

    def test_savings_potential_never_becomes_realized(self):
        docs = self.resolve()
        self.assertIn('savings potential (estimated)', docs.cover_letter_markdown)
        self.assertNotIn('Saved $10,000', docs.cover_letter_markdown)
        self.assertNotIn('realized savings', docs.resume_markdown)
        self.assertIn('savings potential (estimated)', docs.resume_markdown)

    def test_explicit_unit_selection_picks_evidence_without_dumping_whole_passage(self):
        value = fixture()
        value.resume_sections[0].blocks = [value.cover_letter_paragraphs[0]]
        value.resume_sections[0].blocks[0].unit_indices = [2, 6]
        docs = self.resolve(value)
        self.assertIn('Built a workflow tracker.', docs.resume_markdown)
        self.assertIn('Documented insurance verification workflows.', docs.resume_markdown)
        self.assertNotIn('Organized weekly reviews.', docs.resume_markdown)
        self.assertNotIn('Coordinated onboarding for 12 employees.', docs.resume_markdown)

    def test_bad_unit_index_and_unsplittable_giant_bullet_fail(self):
        value = fixture()
        value.cover_letter_paragraphs[0].unit_indices = [999]
        with self.assertRaisesRegex(ValueError, 'invalid source unit'):
            self.resolve(value)
        profile = '## Experience\n- Coordinated ' + ' '.join(['documented'] * 60) + ' workflows.\n'
        value = SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=['P001'], phrasing='source')])],
            cover_letter_paragraphs=[DocumentBlock(evidence_ids=['P001'], phrasing='first_person')], verification_items=[])
        with self.assertRaisesRegex(ValueError, 'exceeds 45 words'):
            self.resolve(value, profile)

    def test_quality_gate_blocks_instruction_leakage_giant_bullets_and_absent_identity(self):
        for text in ['# Jordan Example\n\njordan@example.test\n\n- Do not invent revenue.',
                     '# Jordan Example\n\njordan@example.test\n\n- ' + 'word ' * 46,
                     '# Resume\n\n- Coordinated training.']:
            docs = self.resolve()
            docs.resume_markdown = text
            with self.subTest(text=text[:30]), self.assertRaisesRegex(ValueError, 'quality check'):
                validate_quality(docs)

    def test_separate_employers_under_generic_experience_do_not_mix(self):
        profile = ('## Experience\n- Coordinator at First Example, 2020-2021.\n- Coordinated training.\n'
                   '- Analyst at Second Example, 2022-2023.\n- Analyzed operational reports.\n')
        value = SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=['P002', 'P004'], phrasing='source')])],
            cover_letter_paragraphs=[DocumentBlock(evidence_ids=['P002'], phrasing='first_person')], verification_items=[])
        docs = self.resolve(value, profile)
        self.assertIn('### First Example\n\nCoordinator | 2020-2021\n\n- Coordinated training.', docs.resume_markdown)
        self.assertIn('### Second Example\n\nAnalyst | 2022-2023\n\n- Analyzed operational reports.', docs.resume_markdown)

    def test_explicit_identity_heading_resolves_to_literal_profile_span(self):
        profile = '# Applicant name: Riley Example\n## Contact\n- riley@example.test\n## Experience\n- Coordinated onboarding.\n'
        value = SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=['P002'], phrasing='source')])],
            cover_letter_paragraphs=[DocumentBlock(evidence_ids=['P002'], phrasing='first_person')], verification_items=[])
        docs = self.resolve(value, profile)
        self.assertIn('# Riley Example', docs.resume_markdown)
        entry = next(e for e in docs.evidence if e.claim == 'Riley Example')
        self.assertEqual(profile[entry.source_start:entry.source_end], entry.source_quote)
        self.assertEqual(entry.evidence_id, '')

    def test_linkedin_only_contact_keeps_explicit_email_phone_placeholder(self):
        profile = '## Contact\n- Name: Riley Example\n- LinkedIn: https://linkedin.com/in/riley-example\n## Experience\n- Coordinated onboarding.\n'
        value = SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=['P003'], phrasing='source')])],
            cover_letter_paragraphs=[DocumentBlock(evidence_ids=['P003'], phrasing='first_person')], verification_items=[])
        docs = self.resolve(value, profile)
        self.assertIn('https://linkedin.com/in/riley-example', docs.resume_markdown)
        self.assertIn('[Email / phone — verify]', docs.resume_markdown)

    def test_conflicting_realized_language_is_blocked_instead_of_prefixed(self):
        profile = '## Experience\n- Saved $10,000 in costs. These savings were potential, not realized.\n'
        value = SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=['P001'], unit_indices=[1], phrasing='source')])],
            cover_letter_paragraphs=[DocumentBlock(evidence_ids=['P001'], unit_indices=[1], phrasing='first_person')], verification_items=[])
        with self.assertRaisesRegex(ValueError, 'financial claim conflicts'):
            self.resolve(value, profile)

    def test_desired_role_field_cannot_replace_actual_title(self):
        profile = ('## Experience\n### Example Services\n- Company: Example Services\n'
                   '- Actual job title: Coordinator\n- Dates: 2020-2021\n'
                   '- Desired role: Business Operations Manager\n- Coordinated onboarding.\n')
        value = SourceApplicationDocuments(resume_sections=[ResumeSection(heading='Experience', blocks=[DocumentBlock(evidence_ids=['P005'], phrasing='source')])],
            cover_letter_paragraphs=[DocumentBlock(evidence_ids=['P005'], phrasing='first_person')], verification_items=[])
        docs = self.resolve(value, profile)
        self.assertIn('Coordinator | 2020-2021', docs.resume_markdown)
        self.assertNotIn('Business Operations Manager', docs.resume_markdown)

    def test_contact_and_raw_job_metadata_do_not_become_body_bullets_or_letter_text(self):
        value = fixture()
        value.resume_sections[0].blocks.insert(0, DocumentBlock(evidence_ids=['P001', 'P002'], phrasing='source'))
        company = next(e['evidence_id'] for e in document_catalog(PROFILE) if e['text'].startswith('- Company:'))
        value.cover_letter_paragraphs.insert(0, DocumentBlock(evidence_ids=[company, 'P001'], phrasing='source'))
        docs = self.resolve(value)
        self.assertNotIn('- Name:', docs.resume_markdown)
        self.assertNotIn('Company:', docs.cover_letter_markdown)
        self.assertIn('# Jordan Example', docs.resume_markdown)
        self.assertIn('jordan@example.test', docs.cover_letter_markdown)


if __name__ == '__main__':
    unittest.main()
