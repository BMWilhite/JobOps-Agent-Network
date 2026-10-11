import re
from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(31, 88, 74)


def add_formatted_text(paragraph, text):
    """Convert Markdown bold markers into editable Word formatting."""
    for index, part in enumerate(text.split("**")):
        if not part:
            continue

        run = paragraph.add_run(part)

        if index % 2 == 1:
            run.bold = True

def markdown_to_docx(markdown, destination):
    """Convert simple resume Markdown into an editable Word file."""
    document = Document()
    section = document.sections[0]

    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(5)

    title_style = document.styles["Title"]
    title_style.font.name = "Aptos Display"
    title_style.font.size = Pt(18)
    title_style.font.bold = True
    title_style.font.color.rgb = ACCENT

    for name, size in (("Heading 2", 11), ("Heading 3", 10)):
        style = document.styles[name]
        style.font.name = "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.paragraph_format.space_before = Pt(9)
        style.paragraph_format.space_after = Pt(4)

    for line in markdown.splitlines():
        text = line.strip()

        if not text or text == "---":
            continue

        if text.startswith("# "):
            paragraph = document.add_paragraph(style="Title")
            add_formatted_text(paragraph, text[2:])

        elif text.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_formatted_text(paragraph, text[4:])

        elif text.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_formatted_text(paragraph, text[3:])

        elif text.startswith(("- ", "* ")):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(2)
            add_formatted_text(paragraph, text[2:])

        else:
            paragraph = document.add_paragraph()
            add_formatted_text(paragraph, text)

    destination = Path(destination)
    document.save(destination)
    return destination
